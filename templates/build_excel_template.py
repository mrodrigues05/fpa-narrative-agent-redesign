"""
Template — gera a planilha de dados (dados.xlsx) consumida por
agent_engine_template.py e por uma landing page (via upload + SheetJS).

Isso é um TEMPLATE — preencha os dicionários AJUSTE AQUI com os drivers reais
do seu case antes de rodar. A linha "Margem de contribuição" é escrita como
FÓRMULA (não como valor colado) — depois de rodar este script, recalcule com
o recalc.py do skill de xlsx antes de usar o arquivo em qualquer lugar:

    python3 /mnt/skills/public/xlsx/scripts/recalc.py dados.xlsx

Sem isso, tanto o openpyxl (Python, data_only=True) quanto o SheetJS (no
navegador) leem as células de fórmula como vazias.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

BLUE = Font(name="Arial", color="0000FF")   # inputs (drivers brutos)
BLACK = Font(name="Arial", color="000000")  # fórmulas
BOLD = Font(name="Arial", bold=True)
HEADER = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1C4200")  # troque pela cor da marca do seu case
NUM = "#,##0.00"
PCT = "0.0%"

# ============================================================
# AJUSTE AQUI — dados reais do seu case
# ============================================================
MESES = ["mes1", "mes2", "mes3"]  # ex.: ["jun/26", "jul/26", "ago/26"]

BU_A_LABEL = "Garantia"
BU_A_DADOS = {
    "originacao":     {"mes1": 0, "mes2": 0, "mes3": 0},
    "carteira_media": {"mes1": 0, "mes2": 0, "mes3": 0},
    "carteira_eop":   {"mes1": 0, "mes2": 0, "mes3": 0},
    "taxa":           {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "funding":        {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "credito":        {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "opex":           {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "cac":            {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
}
BU_B_LABEL = "Consignado"
BU_B_DADOS = {
    "originacao":     {"mes1": 0, "mes2": 0, "mes3": 0},
    "carteira_media": {"mes1": 0, "mes2": 0, "mes3": 0},
    "carteira_eop":   {"mes1": 0, "mes2": 0, "mes3": 0},
    "taxa":           {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "funding":        {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "credito":        {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "opex":           {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
    "cac":            {"mes1": 0.0, "mes2": 0.0, "mes3": 0.0},
}
# Se o case fornecer um total de referência para conferir reconciliação:
MARGEM_REFERENCIA = {"mes1": None, "mes2": None, "mes3": None}
FONTE = "Fonte: [cite a origem real do dado — case, planilha do time, sistema]."
# ============================================================


def write_bu_sheet(wb, nome, dados):
    ws = wb.create_sheet(nome)
    ws["A1"] = "Driver"; ws["A1"].font = HEADER; ws["A1"].fill = HEADER_FILL
    for i, mes in enumerate(MESES):
        c = ws.cell(row=1, column=2 + i, value=mes); c.font = HEADER; c.fill = HEADER_FILL

    linhas = [
        ("Originação (R$ mn)", dados["originacao"], NUM, BLUE),
        ("Carteira média (R$ mn)", dados["carteira_media"], NUM, BLUE),
        ("Carteira final - EoP (R$ mn)", dados["carteira_eop"], NUM, BLUE),
        ("Taxa média (% a.a.)", dados["taxa"], PCT, BLUE),
        ("Custo de funding (% a.a.)", dados["funding"], PCT, BLUE),
        ("Custo de crédito (% a.a.)", dados["credito"], PCT, BLUE),
        ("Opex variável (% a.a.)", dados["opex"], PCT, BLUE),
        ("CAC (% da originação)", dados["cac"], PCT, BLUE),
    ]
    row, row_ref = 2, {}
    for label, valores, fmt, font in linhas:
        ws.cell(row=row, column=1, value=label).font = BOLD
        for i, mes in enumerate(MESES):
            cell = ws.cell(row=row, column=2 + i, value=valores[mes])
            cell.number_format = fmt; cell.font = font
        row_ref[label] = row
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="Margem de contribuição (R$ mn)").font = BOLD
    r = {k: row_ref[v] for k, v in {
        "orig": "Originação (R$ mn)", "cart": "Carteira média (R$ mn)", "taxa": "Taxa média (% a.a.)",
        "fund": "Custo de funding (% a.a.)", "cred": "Custo de crédito (% a.a.)", "opex": "Opex variável (% a.a.)",
        "cac": "CAC (% da originação)",
    }.items()}
    for i in range(len(MESES)):
        col = get_column_letter(2 + i)
        formula = (f"={col}{r['cart']}*{col}{r['taxa']}/12-{col}{r['cart']}*{col}{r['fund']}/12"
                   f"-{col}{r['cart']}*{col}{r['cred']}/12-{col}{r['cart']}*{col}{r['opex']}/12"
                   f"-{col}{r['orig']}*{col}{r['cac']}")
        cell = ws.cell(row=row, column=2 + i, value=formula); cell.number_format = NUM; cell.font = BLACK
    margem_row = row
    row += 2
    ws.cell(row=row, column=1, value=(
        "Fórmula: Margem = Carteira média × Taxa a.a. ÷ 12 − Carteira média × Funding ÷ 12 "
        "− Carteira média × Custo de crédito ÷ 12 − Carteira média × Opex ÷ 12 − Originação × CAC."
    )).font = Font(name="Arial", italic=True, size=9, color="666666")
    row += 1
    ws.cell(row=row, column=1, value=FONTE).font = Font(name="Arial", italic=True, size=9, color="666666")

    for col, width in [("A", 30), ("B", 12), ("C", 12), ("D", 12)]:
        ws.column_dimensions[col].width = width
    return margem_row


def build(output_path="dados.xlsx"):
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    mg_row = write_bu_sheet(wb, BU_A_LABEL, BU_A_DADOS)
    mc_row = write_bu_sheet(wb, BU_B_LABEL, BU_B_DADOS)

    ws = wb.create_sheet("Consolidado")
    ws["A1"] = "Indicador"; ws["A1"].font = HEADER; ws["A1"].fill = HEADER_FILL
    for i, mes in enumerate(MESES):
        c = ws.cell(row=1, column=2 + i, value=mes); c.font = HEADER; c.fill = HEADER_FILL
    ws.cell(row=2, column=1, value="Originação total (R$ mn)").font = BOLD
    ws.cell(row=3, column=1, value="Margem de contribuição (R$ mn)").font = BOLD
    ws.cell(row=4, column=1, value="Margem — valor de referência (R$ mn)").font = BOLD
    ws.cell(row=5, column=1, value="Reconciliado? (|calc - referência| < 0,02)").font = BOLD
    for i, mes in enumerate(MESES):
        col = get_column_letter(2 + i)
        ws.cell(row=2, column=2 + i, value=f"={BU_A_LABEL}!{col}2+{BU_B_LABEL}!{col}2").number_format = NUM
        ws.cell(row=3, column=2 + i, value=f"={BU_A_LABEL}!{col}{mg_row}+{BU_B_LABEL}!{col}{mc_row}").number_format = NUM
        if MARGEM_REFERENCIA.get(mes) is not None:
            ws.cell(row=4, column=2 + i, value=MARGEM_REFERENCIA[mes]).number_format = NUM
            ws.cell(row=4, column=2 + i).font = BLUE
            ws.cell(row=5, column=2 + i, value=f'=IF(ABS({col}3-{col}4)<0.02,"OK","DIVERGE")')
        for rr in (2, 3):
            ws.cell(row=rr, column=2 + i).font = BLACK
    for col, width in [("A", 42), ("B", 14), ("C", 14), ("D", 14)]:
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Leia-me", 0)
    for i, l in enumerate([
        "dados.xlsx — drivers por unidade de negócio para o agente FP&A",
        "",
        "Abas por BU: uma linha por driver, uma coluna por mês. A linha 'Margem de",
        "contribuição' é FÓRMULA — recalcula se você mudar qualquer driver acima.",
        "Aba 'Consolidado': soma as BUs e confere reconciliação, se houver referência.",
        "",
        "Não renomeie as abas nem os rótulos da coluna A — o parser casa por texto.",
    ], start=1):
        ws2.cell(row=i, column=1, value=l).font = Font(name="Arial", size=10, bold=(i == 1))
    ws2.column_dimensions["A"].width = 90

    wb.save(output_path)
    print(f"salvo em {output_path} — rode o recalc.py do skill de xlsx antes de usar.")


if __name__ == "__main__":
    build()
