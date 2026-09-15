"""
Template — Agent Engine para redesenho de agente de narrativa financeira (FP&A)
================================================================================

Isso é um TEMPLATE — adapte a seção "AJUSTE AQUI" antes de rodar. O resto
(evidence contract, memória histórica falsificável, checker, harness de 3
casos de teste) é genérico: foi validado no case de referência (Vértice) e
não deveria precisar mudar para um case no mesmo formato (duas unidades de
negócio, um KPI consolidado, um repositório de contexto com um padrão
observado). Ver `motor-fpa/` desta skill para o motor já rodando num case real.

Uso: coloque este arquivo e `dados.xlsx` (ver templates/build_excel_template.py)
na mesma pasta, ajuste a seção abaixo, e rode `python3 agent_engine_template.py`.
"""

import os
import unicodedata
import openpyxl

# ============================================================
# AJUSTE AQUI — o que muda de case para case
# ============================================================
BU_A_LABEL = "Garantia"      # nome da aba no Excel e nome exibido na narrativa
BU_B_LABEL = "Consignado"
BU_A, BU_B = "bu_a", "bu_b"  # chaves internas — não precisam mudar

DRIVER_ORDER = ["carteira_media", "taxa", "funding", "credito", "opex", "originacao", "cac"]


def _calc(vals):
    """A FÓRMULA DO KPI deste case — único ponto do motor com lógica de
    negócio específica do domínio. Troque por completo se o seu KPI não for
    'margem de contribuição de uma carteira de crédito'."""
    return (
        vals["carteira_media"] * vals["taxa"] / 12
        - vals["carteira_media"] * vals["funding"] / 12
        - vals["carteira_media"] * vals["credito"] / 12
        - vals["carteira_media"] * vals["opex"] / 12
        - vals["originacao"] * vals["cac"]
    )


def _match_driver_key(label):
    """Casa o rótulo da coluna A da planilha com uma chave de DRIVER_ORDER.
    Ordem importa: o primeiro casamento vence (por isso 'cac' vem antes de
    'origina' — o rótulo 'CAC (% da originação)' contém as duas). Ajuste as
    palavras-chave se os nomes dos seus drivers forem diferentes."""
    l = "".join(c for c in unicodedata.normalize("NFD", str(label).lower()) if unicodedata.category(c) != "Mn")
    if "cac" in l:
        return "cac"
    if "origina" in l:
        return "originacao"
    if "carteira" in l and "media" in l:
        return "carteira_media"
    if "carteira" in l:
        return None  # provavelmente saldo final (EoP) — informativo, não entra no cálculo
    if "taxa" in l:
        return "taxa"
    if "funding" in l:
        return "funding"
    if "credito" in l:
        return "credito"
    if "opex" in l:
        return "opex"
    return None


# ============================================================
# Daqui para baixo, genérico — normalmente não precisa mexer
# ============================================================
EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dados.xlsx")


def load_drivers_from_excel(path=EXCEL_PATH):
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"DADO AUSENTE: não encontrei '{path}'. O motor não estima no lugar "
            "de um dado que falta — gere/copie a planilha para essa pasta antes de rodar."
        )
    wb = openpyxl.load_workbook(path, data_only=True)
    drivers, meses = {}, None
    for sheet_name, bu_key in [(BU_A_LABEL, BU_A), (BU_B_LABEL, BU_B)]:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"DADO AUSENTE: aba '{sheet_name}' não existe em {path}.")
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        sheet_meses = [m for m in header[1:] if m]
        meses = meses or sheet_meses
        bu_data = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = row[0]
            if not label or len(str(label)) > 50:
                continue
            key = _match_driver_key(label)
            if key is None:
                continue
            bu_data[key] = {mes: row[1 + i] for i, mes in enumerate(sheet_meses) if row[1 + i] is not None}
        for k in DRIVER_ORDER:
            if k not in bu_data:
                raise ValueError(f"DADO AUSENTE: driver '{k}' não encontrado na aba '{sheet_name}'.")
        drivers[bu_key] = bu_data
    return drivers, meses


DRIVERS, MESES_REAIS = load_drivers_from_excel()
print(f"[agent_engine_template] drivers carregados de {EXCEL_PATH} — meses: {MESES_REAIS}")


def add_synthetic_scenario(drivers, base_month, target_delta=1.0):
    """Cenário SINTÉTICO para o Caso C (evidência insuficiente) — não é dado
    real. Desloca a taxa de cada BU o suficiente para gerar um Δ de
    magnitude parecida e sinal oposto entre as duas unidades."""
    for bu, sinal in [(BU_A, -1), (BU_B, 1)]:
        d = drivers[bu]
        for k in DRIVER_ORDER:
            if k != "taxa":
                d[k]["set_sintetico"] = d[k][base_month]
        delta = sinal * target_delta * 12 / d["carteira_media"][base_month]
        d["taxa"]["set_sintetico"] = d["taxa"][base_month] + delta


add_synthetic_scenario(DRIVERS, MESES_REAIS[-1])


def margem(bu, mes):
    d = DRIVERS[bu]
    return _calc({k: d[k][mes] for k in DRIVER_ORDER})


def bridge_bu(bu, mes_a, mes_b):
    d = DRIVERS[bu]
    cur = {k: d[k][mes_a] for k in DRIVER_ORDER}
    efeitos = {}
    for drv in DRIVER_ORDER:
        novo = dict(cur)
        novo[drv] = d[drv][mes_b]
        efeitos[drv] = round(_calc(novo) - _calc(cur), 3)
        cur = novo
    return efeitos


def build_evidence(mes_atual, mes_comp):
    mg_atual, mg_comp = margem(BU_A, mes_atual), margem(BU_A, mes_comp)
    mc_atual, mc_comp = margem(BU_B, mes_atual), margem(BU_B, mes_comp)
    dg, dc = mg_atual - mg_comp, mc_atual - mc_comp
    consolidado_atual, consolidado_comp = mg_atual + mc_atual, mg_comp + mc_comp
    delta_consolidado = consolidado_atual - consolidado_comp
    bridge_sum = dg + dc
    reconciled = abs(bridge_sum - delta_consolidado) < 0.02

    rel_gap = abs(abs(dg) - abs(dc)) / max(abs(dg), abs(dc), 1e-9)
    if rel_gap < 0.15:
        primary = "INSUFFICIENT_EVIDENCE"
    else:
        primary = BU_A_LABEL if abs(dg) > abs(dc) else BU_B_LABEL

    drivers_detail = None
    if primary == BU_A_LABEL:
        drivers_detail = bridge_bu(BU_A, mes_comp, mes_atual)
    elif primary == BU_B_LABEL:
        drivers_detail = bridge_bu(BU_B, mes_comp, mes_atual)

    return {
        "period": mes_atual,
        "comparison_period": mes_comp,
        "margin": {
            "current": round(consolidado_atual, 2),
            "previous": round(consolidado_comp, 2),
            "delta": round(delta_consolidado, 2),
        },
        "bu_contributions": [
            {"bu": BU_A_LABEL, "delta": round(dg, 2)},
            {"bu": BU_B_LABEL, "delta": round(dc, 2)},
        ],
        "primary_driver": primary,
        "reconciliation": {"bridge_sum": round(bridge_sum, 2), "reconciled": reconciled},
        "drivers_detail": drivers_detail,
    }


# ---------- memória histórica: claims falsificáveis, não regras permanentes ----------
# AJUSTE AQUI — descreva a claim que aparece no repositório de contexto do seu case.
HISTORICAL_MEMORY = [
    {
        "claim_id": "TROQUE_PELO_ID_DA_CLAIM_DO_SEU_CASE",
        "descricao": "Descrição em uma frase da hipótese causal que o repositório registrou.",
        "period_origem": "TROQUE",
        "evidence_origem": "TROQUE",
        "valid_until": "requires_current_period_validation",
        "bu_relacionada": BU_B_LABEL,  # a unidade que a claim aponta como causa
    }
]


def validate_hypothesis(claim, evidence):
    """Nunca aceita o histórico como prova — testa de novo contra o período atual."""
    bu_delta = next(b["delta"] for b in evidence["bu_contributions"] if b["bu"] == claim["bu_relacionada"])
    supported = bu_delta < 0 and evidence["primary_driver"] == claim["bu_relacionada"]
    status = "confirmed_for_period" if supported else "contradicted"
    return {
        **claim,
        "period_testado": evidence["period"],
        "status": status,
        "evidence_testada": (
            f"delta({claim['bu_relacionada']}, {evidence['period']})={bu_delta:+.2f}; "
            f"primary_driver={evidence['primary_driver']}"
        ),
    }


def render_reference_narrative(evidence, hypotheses_validadas):
    if evidence["primary_driver"] == "INSUFFICIENT_EVIDENCE":
        a, b = evidence["bu_contributions"]
        return (
            f"Δ consolidado de {evidence['margin']['delta']:+.2f} entre "
            f"{evidence['comparison_period']} e {evidence['period']}. {a['bu']} ({a['delta']:+.2f}) e "
            f"{b['bu']} ({b['delta']:+.2f}) têm impacto de magnitude parecida e sinal oposto. Não há "
            "evidência suficiente, com os dados disponíveis, para apontar um driver principal isolado."
        )
    contradita = any(h["status"] == "contradicted" for h in hypotheses_validadas)
    aviso = " A hipótese histórica não se sustenta neste período e foi rejeitada." if contradita else ""
    dom = next(b for b in evidence["bu_contributions"] if b["bu"] == evidence["primary_driver"])
    outro = next(b for b in evidence["bu_contributions"] if b["bu"] != evidence["primary_driver"])
    return (
        f"A margem consolidada foi de {evidence['margin']['current']:.2f} em {evidence['period']}, "
        f"variação de {evidence['margin']['delta']:+.2f} contra {evidence['comparison_period']}.{aviso} "
        f"O driver principal foi {evidence['primary_driver']} (Δ {dom['delta']:+.2f}); {outro['bu']} teve "
        f"Δ {outro['delta']:+.2f} no período."
    )


def run_checker(narrativa_texto, evidence, hypotheses_validadas):
    msgs = []
    reconciled = evidence["reconciliation"]["reconciled"]
    msgs.append(("Reconciliação", reconciled,
                 f"bridge_sum={evidence['reconciliation']['bridge_sum']:+.2f} vs "
                 f"delta_consolidado={evidence['margin']['delta']:+.2f}"))

    primary = evidence["primary_driver"]
    if primary == "INSUFFICIENT_EVIDENCE":
        primary_ok = "evidência suficiente" in narrativa_texto.lower()
        msgs.append(("Primary driver", primary_ok,
                     "evidência é ambígua; narrativa deve declarar limitação, não escolher um driver"))
    else:
        primary_ok = primary.lower() in narrativa_texto.lower()
        msgs.append(("Primary driver", primary_ok, f"narrativa deve citar '{primary}' como driver principal"))

    anchoring_detected = False
    for h in hypotheses_validadas:
        if h["status"] == "contradicted":
            t = narrativa_texto.lower()
            if "não se sustenta" not in t and "rejeitada" not in t:
                anchoring_detected = True
    msgs.append(("Historical anchoring", not anchoring_detected,
                 "narrativa não pode reafirmar uma hipótese já contradita pelo período atual"))

    passed = [ok for _, ok, _ in msgs]
    if all(passed):
        confidence = "HIGH" if primary != "INSUFFICIENT_EVIDENCE" else "MEDIUM"
    elif reconciled:
        confidence = "LOW"
    else:
        confidence = "BLOCKED"
    return {"status": "PASS" if all(passed) else "FAIL", "confidence": confidence, "checks": msgs}


def rodar_caso(nome, mes_atual, mes_comp, usar_memoria=True):
    print(f"\n=== Caso {nome}: {mes_comp} -> {mes_atual} ===")
    evidence = build_evidence(mes_atual, mes_comp)
    hyps = [validate_hypothesis(c, evidence) for c in HISTORICAL_MEMORY] if usar_memoria else []
    for h in hyps:
        print(f"  [historical hypothesis] status={h['status']} ({h['evidence_testada']})")
    narrativa = render_reference_narrative(evidence, hyps)
    print(f"  narrativa: {narrativa}")
    resultado = run_checker(narrativa, evidence, hyps)
    for nome_check, ok, detalhe in resultado["checks"]:
        print(f"  [{'OK' if ok else 'FALHOU'}] {nome_check}: {detalhe}")
    print(f"  >> checker status={resultado['status']}  confidence={resultado['confidence']}")
    return evidence, hyps, narrativa, resultado


if __name__ == "__main__":
    # AJUSTE AQUI — troque pelos períodos reais do seu case.
    _, hyps_a, _, res_a = rodar_caso("A (padrão confirmado)", MESES_REAIS[1], MESES_REAIS[0])
    ev_b, hyps_b, narr_b, res_b = rodar_caso("B (adversarial)", MESES_REAIS[2], MESES_REAIS[1])
    _, _, _, res_c = rodar_caso("C (evidência insuficiente, sintético)", "set_sintetico", MESES_REAIS[-1], usar_memoria=False)

    print("\n=== Teste de regressão (Historical Anchoring Test) ===")
    assert hyps_a[0]["status"] == "confirmed_for_period", "Caso A deveria confirmar a hipótese"
    assert hyps_b[0]["status"] == "contradicted", "Caso B deveria contradizer a hipótese"
    assert res_b["status"] == "PASS", "Caso B: narrativa de referência deveria passar no checker"
    assert res_c["confidence"] in ("MEDIUM", "LOW"), "Caso C deveria sinalizar confiança reduzida"
    print("OK — adapte HISTORICAL_MEMORY e os períodos acima ao case real antes de usar isso como prova.")
