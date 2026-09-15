"""
Protótipo v2 — Agent Engine (DADO -> CÁLCULO -> EVIDÊNCIA -> MEMÓRIA -> CHECKER)
=================================================================================

Esta é a camada DETERMINÍSTICA do agente redesenhado. Ela nunca escreve
narrativa — produz um contrato de evidência estruturado (mesmo shape que a
landing page `agente_fpa_landing.html` consome em JS) e valida qualquer
narrativa recebida contra esse contrato.

Correspondência com a landing page:
  build_evidence()              -> painel "Evidence"
  HISTORICAL_MEMORY / validate_hypothesis() -> aba "Historical Context"
  render_reference_narrative()  -> modo "Determinístico (offline)" do seletor de execução
  run_checker()                 -> painel "Validation" / Narrative Check

Por que o template de narrativa aqui é determinístico, não um LLM real: este
sandbox não tem acesso à rede (ver network_configuration), então a chamada
real à API não pode ser testada a partir do Python. Na landing page, o mesmo
motor roda em JS e o passo de narrativa tenta uma chamada real à API da
Anthropic, com este template como fallback "modo offline". O ponto central do
redesenho não muda: o LLM nunca decide sozinho o primary_driver — ele só
recebe a evidência já calculada.

DADOS: os drivers de jun/jul/ago vêm de `margem.xlsx` (abas Garantia e
Consignado), não de números embutidos no código — é o "motor de consolidação
exporta um CSV/planilha mensal" do pipeline original, só que de verdade. Se o
arquivo não existir ao lado deste script, o motor falha alto e explica o que
falta, em vez de usar um número antigo silenciosamente (regra §16 do redesenho:
"dado ausente" nunca vira "estimativa plausível").
"""

import os
import openpyxl

DRIVER_ORDER = ["carteira_media", "taxa", "funding", "credito", "opex", "originacao", "cac"]
EXCEL_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "margem.xlsx")


def _match_driver_key(label):
    """Casa o rótulo da coluna A da planilha com a chave interna do motor.
    Casamento por palavra-chave, não por posição — a ordem das linhas na
    planilha pode mudar sem quebrar o parser."""
    l = str(label).lower()
    if "cac" in l:
        return "cac"
    if "origina" in l:
        return "originacao"
    if "carteira" in l and ("média" in l or "media" in l):
        return "carteira_media"
    if "carteira" in l and ("eop" in l or "final" in l):
        return None  # informativo — não entra na fórmula de margem
    if "taxa" in l:
        return "taxa"
    if "funding" in l:
        return "funding"
    if "cr" in l and ("dito" in l or "édito" in l):
        return "credito"
    if "opex" in l:
        return "opex"
    return None


def load_drivers_from_excel(path=EXCEL_PATH):
    if not os.path.exists(path):
        raise FileNotFoundError(
            f"DADO AUSENTE: não encontrei '{path}'. O motor não estima no lugar "
            "de um dado que falta — gere/copie margem.xlsx para essa pasta antes de rodar."
        )
    wb = openpyxl.load_workbook(path, data_only=True)
    drivers, meses = {}, None
    for sheet_name, bu_key in [("Garantia", "garantia"), ("Consignado", "consignado")]:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"DADO AUSENTE: aba '{sheet_name}' não existe em {path}.")
        ws = wb[sheet_name]
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        sheet_meses = [m for m in header[1:] if m]
        meses = meses or sheet_meses
        bu_data = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = row[0]
            if not label or len(str(label)) > 50:
                continue  # ignora linhas em branco e as notas de rodapé (texto longo)
            key = _match_driver_key(label)
            if key is None:
                continue
            bu_data[key] = {mes: row[1 + i] for i, mes in enumerate(sheet_meses) if row[1 + i] is not None}
        for k in ["originacao", "carteira_media", "taxa", "funding", "credito", "opex", "cac"]:
            if k not in bu_data:
                raise ValueError(f"DADO AUSENTE: driver '{k}' não encontrado na aba '{sheet_name}'.")
        drivers[bu_key] = bu_data
    return drivers, meses


DRIVERS, MESES_REAIS = load_drivers_from_excel()
print(f"[agent_engine] drivers carregados de {EXCEL_PATH} — meses: {MESES_REAIS}")

# Cenário SINTÉTICO — não vem da planilha nem do case. Existe só para exercitar
# o caminho de código "evidência insuficiente" (Caso C, doc de redesenho §14):
# desloca a taxa de cada BU o suficiente para gerar ~R$1mn de Δ margem em
# sinais opostos, calculado a partir do último mês carregado — funciona com
# qualquer planilha, não só com jun/jul/ago do case.
def add_synthetic_scenario(drivers, base_month, target_delta=1.0):
    for bu, sinal in [("garantia", -1), ("consignado", +1)]:
        d = drivers[bu]
        for k in ["originacao", "carteira_media", "funding", "credito", "opex", "cac"]:
            d[k]["set_sintetico"] = d[k][base_month]
        delta_taxa = sinal * target_delta * 12 / d["carteira_media"][base_month]
        d["taxa"]["set_sintetico"] = d["taxa"][base_month] + delta_taxa


add_synthetic_scenario(DRIVERS, MESES_REAIS[-1])


def _calc(vals):
    return (
        vals["carteira_media"] * vals["taxa"] / 12
        - vals["carteira_media"] * vals["funding"] / 12
        - vals["carteira_media"] * vals["credito"] / 12
        - vals["carteira_media"] * vals["opex"] / 12
        - vals["originacao"] * vals["cac"]
    )


def margem(bu, mes):
    d = DRIVERS[bu]
    return _calc({k: d[k][mes] for k in DRIVER_ORDER})


def bridge_bu(bu, mes_a, mes_b):
    d = DRIVERS[bu]
    cur = {k: d[k][mes_a] for k in DRIVER_ORDER}
    efeitos = {}
    for drv in DRIVER_ORDER:
        novo = dict(cur)
        novo[drv] = d[drv][mes_b]
        efeitos[drv] = round(_calc(novo) - _calc(cur), 3)
        cur = novo
    return efeitos


def mix_consignado(mes):
    tot = DRIVERS["garantia"]["originacao"][mes] + DRIVERS["consignado"]["originacao"][mes]
    return DRIVERS["consignado"]["originacao"][mes] / tot


def build_evidence(mes_atual, mes_comp):
    """Contrato de evidência — é isto, e só isto, que o LLM recebe para narrar."""
    mg_atual, mg_comp = margem("garantia", mes_atual), margem("garantia", mes_comp)
    mc_atual, mc_comp = margem("consignado", mes_atual), margem("consignado", mes_comp)
    dg, dc = mg_atual - mg_comp, mc_atual - mc_comp
    consolidado_atual, consolidado_comp = mg_atual + mc_atual, mg_comp + mc_comp
    delta_consolidado = consolidado_atual - consolidado_comp
    bridge_sum = dg + dc
    reconciled = abs(bridge_sum - delta_consolidado) < 0.02

    rel_gap = abs(abs(dg) - abs(dc)) / max(abs(dg), abs(dc), 1e-9)
    if rel_gap < 0.15:
        primary = "INSUFFICIENT_EVIDENCE"
    else:
        primary = "Garantia" if abs(dg) > abs(dc) else "Consignado"

    drivers_detail = None
    if primary in ("Garantia", "Consignado"):
        drivers_detail = bridge_bu(primary.lower(), mes_comp, mes_atual)

    return {
        "period": mes_atual,
        "comparison_period": mes_comp,
        "margin": {
            "current": round(consolidado_atual, 2),
            "previous": round(consolidado_comp, 2),
            "delta": round(delta_consolidado, 2),
        },
        "bu_contributions": [
            {"bu": "Garantia", "delta": round(dg, 2)},
            {"bu": "Consignado", "delta": round(dc, 2)},
        ],
        "primary_driver": primary,
        "reconciliation": {"bridge_sum": round(bridge_sum, 2), "reconciled": reconciled},
        "mix": {
            "consignado_share_comp": round(mix_consignado(mes_comp), 3),
            "consignado_share_atual": round(mix_consignado(mes_atual), 3),
        },
        "drivers_detail": drivers_detail,
    }


# ---------- memória histórica: claims falsificáveis, não regras permanentes ----------
HISTORICAL_MEMORY = [
    {
        "claim_id": "mix_consignado_pressiona_margem",
        "descricao": (
            "Aumento da participação do Consignado na originação pressiona a "
            "margem consolidada (CAC do Consignado é ~2,6x o de Garantia)."
        ),
        "period_origem": "jul/26",
        "evidence_origem": "delta_margin_consignado(jul/26)=-1.87; primary_driver=Consignado",
        "valid_until": "requires_current_period_validation",
    }
]


def validate_hypothesis(claim, evidence):
    """Nunca aceita o histórico como prova — testa de novo contra o período atual."""
    if claim["claim_id"] != "mix_consignado_pressiona_margem":
        return {**claim, "status": "unknown_claim_type"}
    consignado_delta = next(b["delta"] for b in evidence["bu_contributions"] if b["bu"] == "Consignado")
    supported = consignado_delta < 0 and evidence["primary_driver"] == "Consignado"
    status = "confirmed_for_period" if supported else "contradicted"
    return {
        **claim,
        "period_testado": evidence["period"],
        "status": status,
        "evidence_testada": (
            f"delta_margin_consignado({evidence['period']})={consignado_delta:+.2f}; "
            f"primary_driver={evidence['primary_driver']}"
        ),
    }


# ---------- narrativa de referência (determinística, modo offline) ----------
def render_reference_narrative(evidence, hypotheses_validadas):
    if evidence["primary_driver"] == "INSUFFICIENT_EVIDENCE":
        return (
            f"Δ consolidado de {evidence['margin']['delta']:+.2f} mn entre "
            f"{evidence['comparison_period']} e {evidence['period']}. Garantia "
            f"({evidence['bu_contributions'][0]['delta']:+.2f}) e Consignado "
            f"({evidence['bu_contributions'][1]['delta']:+.2f}) têm impacto de magnitude "
            "parecida e sinal oposto. Não há evidência suficiente, com os dados "
            "disponíveis, para apontar um driver principal isolado."
        )
    contradita = any(h["status"] == "contradicted" for h in hypotheses_validadas)
    aviso_hist = (
        " A hipótese histórica de mix não se sustenta neste período e foi rejeitada."
        if contradita else ""
    )
    dom = next(b for b in evidence["bu_contributions"] if b["bu"] == evidence["primary_driver"])
    outro = next(b for b in evidence["bu_contributions"] if b["bu"] != evidence["primary_driver"])
    return (
        f"A margem consolidada foi de R$ {evidence['margin']['current']:.2f} mn em "
        f"{evidence['period']}, variação de R$ {evidence['margin']['delta']:+.2f} mn "
        f"contra {evidence['comparison_period']}.{aviso_hist} O driver principal foi "
        f"{evidence['primary_driver']} (Δ {dom['delta']:+.2f} mn); {outro['bu']} teve "
        f"Δ {outro['delta']:+.2f} mn no período."
    )


# ---------- narrative checker ----------
def run_checker(narrativa_texto, evidence, hypotheses_validadas):
    msgs = []
    reconciled = evidence["reconciliation"]["reconciled"]
    msgs.append(("Reconciliação", reconciled,
                 f"bridge_sum={evidence['reconciliation']['bridge_sum']:+.2f} vs "
                 f"delta_consolidado={evidence['margin']['delta']:+.2f}"))

    primary = evidence["primary_driver"]
    if primary == "INSUFFICIENT_EVIDENCE":
        primary_ok = "não há evidência suficiente" in narrativa_texto.lower() or \
                     "evidência suficiente" in narrativa_texto.lower()
        msgs.append(("Primary driver", primary_ok,
                     "evidência é ambígua; narrativa deve declarar limitação, não escolher um driver"))
    else:
        primary_ok = primary.lower() in narrativa_texto.lower()
        msgs.append(("Primary driver", primary_ok,
                     f"narrativa deve citar '{primary}' como driver principal"))

    anchoring_detected = False
    for h in hypotheses_validadas:
        if h["status"] == "contradicted":
            menciona_mix_sem_rejeitar = (
                "mix" in narrativa_texto.lower() and "não se sustenta" not in narrativa_texto.lower()
                and "rejeitada" not in narrativa_texto.lower() and "não veio do mix" not in narrativa_texto.lower()
                and "não foi" not in narrativa_texto.lower()
            )
            if menciona_mix_sem_rejeitar:
                anchoring_detected = True
    msgs.append(("Historical anchoring", not anchoring_detected,
                 "narrativa não pode reafirmar uma hipótese já contradita pelo período atual"))

    passed = [ok for _, ok, _ in msgs]
    if all(passed):
        confidence = "HIGH" if primary != "INSUFFICIENT_EVIDENCE" else "MEDIUM"
    elif reconciled:
        confidence = "LOW"
    else:
        confidence = "BLOCKED"

    return {
        "status": "PASS" if all(passed) else "FAIL",
        "confidence": confidence,
        "checks": msgs,
    }


def rodar_caso(nome, mes_atual, mes_comp, usar_memoria=True):
    print(f"\n=== Caso {nome}: {mes_comp} -> {mes_atual} ===")
    evidence = build_evidence(mes_atual, mes_comp)
    hyps = [validate_hypothesis(c, evidence) for c in HISTORICAL_MEMORY] if usar_memoria else []
    for h in hyps:
        print(f"  [historical hypothesis] status={h['status']} ({h['evidence_testada']})")
    narrativa = render_reference_narrative(evidence, hyps)
    print(f"  narrativa: {narrativa}")
    resultado = run_checker(narrativa, evidence, hyps)
    for nome_check, ok, detalhe in resultado["checks"]:
        print(f"  [{'OK' if ok else 'FALHOU'}] {nome_check}: {detalhe}")
    print(f"  >> checker status={resultado['status']}  confidence={resultado['confidence']}")
    return evidence, hyps, narrativa, resultado


if __name__ == "__main__":
    # Caso A — padrão histórico confirmado (jun -> jul): mix sobe, Consignado piora.
    _, hyps_a, _, res_a = rodar_caso("A (padrão confirmado)", "jul/26", "jun/26")

    # Caso B — padrão histórico quebrado / ADVERSARIAL (jul -> ago): mix cai,
    # Consignado melhora, Garantia domina. É o caso real do case.
    ev_b, hyps_b, narr_b, res_b = rodar_caso("B (adversarial — mix quebrado)", "ago/26", "jul/26")

    # Caso C — evidência insuficiente (cenário SINTÉTICO, não é dado do case).
    _, _, _, res_c = rodar_caso("C (evidência insuficiente, cenário sintético)",
                                 "set_sintetico", "ago/26", usar_memoria=False)

    print("\n=== Teste de regressão (Historical Anchoring Test) ===")
    assert hyps_a[0]["status"] == "confirmed_for_period", "Caso A deveria confirmar a hipótese"
    assert hyps_b[0]["status"] == "contradicted", "Caso B deveria contradizer a hipótese"
    assert ev_b["primary_driver"] == "Garantia", "Caso B: driver principal deveria ser Garantia"
    assert res_b["status"] == "PASS", "Caso B: narrativa de referência deveria passar no checker"
    assert res_c["confidence"] in ("MEDIUM", "LOW"), "Caso C deveria sinalizar confiança reduzida"
    print("OK — hipótese confirmada em A, contradita em B, driver correto identificado,")
    print("     e evidência ambígua do Caso C não vira um driver inventado.")
